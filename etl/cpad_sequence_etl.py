"""
PRISE — CPAD 2.0 sequence & structure ETL
=========================================

Build-time ETL for CPAD 2.0's **sequence layer** (the non-kinetics workbooks),
producing per-UniProt lookups that power M6's **sequence axis** and
**structure-linkage** panel (PRISE_DESIGN.md §1.1, §3-M6):

  * `aggregating peptides.xlsx`  → per-protein predictor scores
        TANGO, AGGRESCAN, PASTA 2.0, and Waltz-DB/AmyLoad source tags
  * `APR information.xlsx`       → curated aggregation-prone regions per protein
  * `amyloid structure.xlsx`     → CPAD's amyloid structures (PDB cross-refs)

Faithfulness rules (mirroring the kinetics ETL):
  * load/structure/tag only — no scoring or judgement (that is M6).
  * the predictor scores are taken **as bundled in CPAD 2.0**; their individual
    upstream tool versions are NOT recorded by CPAD, so we tag them
    `version: "as-bundled-in-CPAD-2.0"` and never imply a specific tool release.
  * **Zyggregator and CamSol are NOT present** in this download — M6 declares them
    explicitly absent rather than silently dropping the sub-axes.

Each predictor is reduced to one comparable-to-kinetics number with an EXPLICIT
reduction rule (design: "the reduction used to compare to kinetics is explicit per
predictor"); the per-peptide values are summarised, not invented.

Caveats baked into the reductions:
  * predictor pools are collected over WILD-TYPE peptides only (the 'Mutant' column);
    engineered mutants are excluded so a designed destabiliser cannot drag a
    protein's "native" score. Proteins with only mutant peptides emit a record with
    predictors unavailable ("only mutant peptides deposited").
  * the AGGRESCAN/TANGO reductions are FRAGMENT-LEVEL maxima over the deposited WT
    peptides, NOT a whole-sequence native scan — they are a lower bound on the true
    native maximum (documented limitation, not corrected here).
  * PASTA 2.0's +10000 "no cross-β pairing" sentinel is treated as MISSING before
    both the min-reduction and the summary.
  * Waltz is reported as COVERAGE (count of validated amyloid hexapeptides), not as
    a propensity magnitude; absence means UNTESTED, not non-amyloid.

Usage:
    python etl/cpad_sequence_etl.py [--rawdir data/cpad_raw] [--outdir data/processed]

Output (under --outdir):
    sequence_structure.json   per-UniProt {predictors, aprs, structures}
    sequence_manifest.json    version + counts + source SHA-256s
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
from collections import defaultdict
from pathlib import Path
from typing import Any

import openpyxl

SCHEMA_VERSION = "cpad-sequence-1.1"   # v1.1: PASTA sentinel + WT-only reductions (M6 review)
PREDICTOR_VERSION = "as-bundled-in-CPAD-2.0"

# PASTA 2.0 uses +10000 as a sentinel for "no cross-β pairing found" (NOT a real,
# extremely *unfavourable* energy). It must be treated as MISSING before any
# reduction OR summary, otherwise it poisons the min() (fine) AND the mean (e.g.
# Aβ/P05067 showed a fake +103 mean instead of ~ -4). See review item C4.
_PASTA_NO_PAIRING_SENTINEL = 10000.0

# Wild-type marker vocabulary in the peptide sheet's 'Mutant' column. Anything
# else is an engineered point/multi-mutant we must NOT fold into the WT reduction
# (a designed destabiliser would otherwise drag the protein's "native" score).
_WT_MUTANT_TOKENS = {None, "", "no", "wt", "wild type", "wildtype"}


# ------------------------------- helpers ----------------------------------- #
def _clean(v: Any):
    if v is None:
        return None
    s = str(v).strip()
    if s == "" or s.lower() in ("none", "n/a", "na", "nan", "no structures"):
        return None
    return s


def _num(v: Any):
    s = _clean(v)
    if s is None:
        return None
    # Accept integers, decimals AND scientific notation (e.g. "1.2e-3", "4E+05").
    # The exponent is optional; a bare "1e" never matches because [eE] requires
    # the trailing digits. (Review item 15.)
    m = re.search(r"-?\d+(?:\.\d+)?(?:[eE][+-]?\d+)?", s)
    return float(m.group(0)) if m else None


def _is_wt(mutant_raw: Any) -> bool:
    """True when this peptide row is a wild-type (non-engineered) construct.
    Excludes any populated mutation code (e.g. 'A30G', 'Q686K, A692P')."""
    s = _clean(mutant_raw)
    return (s.strip().lower() if s else None) in _WT_MUTANT_TOKENS


def _drop_pasta_sentinel(vals: list) -> list:
    """Strip PASTA 'no cross-β pairing' sentinels (>= 10000) -> treat as MISSING."""
    return [v for v in vals if v is not None and v < _PASTA_NO_PAIRING_SENTINEL]


def sha256_of(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def _header_index(ws):
    hdr = [str(c) if c is not None else "" for c in
           next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    return {name.strip(): i for i, name in enumerate(hdr)}


def _summary(vals: list[float]) -> dict | None:
    vals = [v for v in vals if v is not None]
    if not vals:
        return None
    return {"n": len(vals), "max": max(vals), "min": min(vals),
            "mean": round(sum(vals) / len(vals), 4)}


# ------------------------------ peptide sheet ------------------------------ #
def parse_peptides(path: Path) -> dict:
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb["peptide"]
    h = _header_index(ws)

    def col(name):
        return h.get(name)

    iu = col("Uniprot ID")
    cls = col("Classification")
    tan = col("Tango")
    agg = col("Normalized Aggregation Propensity (AGGRESCAN)")
    energy = col("Best Energy Score (PASTA 2.0)")
    src = col("Source (Waltz-DB,CPAD,AmyLoad,Waltz)")
    name_c = col("Protein Name")
    mutant_c = col("Mutant")
    length_c = col("Length")
    nuapr_c = col("NuAPRpred")          # populated but intentionally excluded (see below)

    by_up: dict[str, dict] = defaultdict(lambda: {
        "protein_name": None, "n_peptides": 0, "n_amyloid": 0, "n_wt": 0,
        # Predictor pools are collected on WILD-TYPE rows ONLY (review item C5):
        # engineered mutants are designed perturbations, not the native sequence,
        # so folding them into a whole-protein reduction is a category error.
        "_tango": [], "_aggrescan": [], "_pasta_energy": [],
        "n_waltz_amyloid": 0, "_lengths": [], "_has_nuapr": False})
    for r in ws.iter_rows(min_row=2, values_only=True):
        up = _clean(r[iu]) if iu is not None else None
        if not up:
            continue
        d = by_up[up]
        d["n_peptides"] += 1
        if d["protein_name"] is None and name_c is not None:
            d["protein_name"] = _clean(r[name_c])
        classification = (_clean(r[cls]) if cls is not None else None) or ""
        is_amyloid = classification.lower() == "amyloid"
        if is_amyloid:
            d["n_amyloid"] += 1
        # NuAPRpred is present in the sheet but deliberately NOT surfaced as a
        # sub-axis (no endpoint contract / reduction defined for it yet); we only
        # record that it *exists* so its exclusion is auditable (review item 14).
        if nuapr_c is not None and _num(r[nuapr_c]) is not None:
            d["_has_nuapr"] = True

        is_wt = _is_wt(r[mutant_c]) if mutant_c is not None else True
        if not is_wt:
            continue                    # mutants excluded from ALL WT reductions
        d["n_wt"] += 1
        if length_c is not None:
            ln = _num(r[length_c])
            if ln is not None:
                d["_lengths"].append(ln)
        for key, idx in (("_tango", tan), ("_aggrescan", agg),
                         ("_pasta_energy", energy)):
            if idx is not None:
                v = _num(r[idx])
                if v is not None:
                    d[key].append(v)
        source = (_clean(r[src]) if src is not None else None) or ""
        if is_amyloid and "waltz" in source.lower():
            d["n_waltz_amyloid"] += 1

    out = {}
    for up, d in by_up.items():
        # PASTA sentinel must be dropped BEFORE both the reduction and the summary.
        pasta_clean = _drop_pasta_sentinel(d["_pasta_energy"])
        lengths = d["_lengths"]
        length_range = [int(min(lengths)), int(max(lengths))] if lengths else None
        excluded = ["NuAPRpred"] if d["_has_nuapr"] else []

        if d["n_wt"] == 0:
            # Zero-WT proteins still emit a record, but predictors are unavailable:
            # we have only engineered mutants, never the native sequence (item C5).
            out[up] = {
                "protein_name": d["protein_name"],
                "n_peptides": d["n_peptides"],
                "n_peptides_total": d["n_peptides"],
                "n_peptides_wt": 0,
                "n_amyloid_peptides": d["n_amyloid"],
                "peptide_length_range": None,
                "excluded_available_predictors": excluded,
                "predictors": _no_wt_predictors(),
            }
            continue

        out[up] = {
            "protein_name": d["protein_name"],
            "n_peptides": d["n_peptides"],            # back-compat alias = total
            "n_peptides_total": d["n_peptides"],
            "n_peptides_wt": d["n_wt"],
            "n_amyloid_peptides": d["n_amyloid"],
            "peptide_length_range": length_range,     # min..max over WT peptides
            # NuAPRpred column is populated but intentionally not exposed as a
            # sub-axis — recorded here so the omission is auditable, not silent.
            "excluded_available_predictors": excluded,
            # de-conflated named sub-axes — each with granularity, endpoint, version,
            # and the EXPLICIT reduction used to compare to a kinetic observable.
            # NOTE: TANGO/AGGRESCAN reductions are FRAGMENT-LEVEL maxima over the
            # deposited WT peptides, NOT a whole-sequence native scan — they are a
            # lower bound on the native maximum (deferred item, documented only).
            "predictors": {
                "tango": _sub_axis(
                    _reduce(d["_tango"], "max"), "residue/segment",
                    "beta-aggregation nucleation",
                    "max segment TANGO score over WT peptides",
                    _summary(d["_tango"])),
                "aggrescan": _sub_axis(
                    _reduce(d["_aggrescan"], "max"), "residue",
                    "generic-aggregation",
                    "max normalized aggregation propensity over WT peptides",
                    _summary(d["_aggrescan"])),
                "pasta": _sub_axis(
                    _reduce(pasta_clean, "min"), "residue-pair",
                    "amyloid (cross-beta)",
                    "min (most stable) PASTA pairing energy over WT peptides "
                    "(no-pairing sentinel +10000 treated as missing)",
                    _summary(pasta_clean)),
                # Waltz here is COVERAGE, not a propensity magnitude: the raw count
                # of Waltz-validated amyloid-positive hexapeptides. Presenting that
                # count as a sub-axis "value" implied an intensity it does not have,
                # so value=None and availability = "has a tested hexapeptide"
                # (coverage). Absence therefore means UNTESTED, not non-amyloid.
                "waltz": _waltz_sub_axis(d["n_waltz_amyloid"]),
            },
        }
    return out


def _reduce(vals, how):
    vals = [v for v in vals if v is not None]
    if not vals:
        return None
    return max(vals) if how == "max" else min(vals)


def _sub_axis(value, granularity, endpoint, reduction, summary):
    return {"value": value, "granularity": granularity, "endpoint": endpoint,
            "reduction": reduction, "version": PREDICTOR_VERSION,
            "summary": summary,
            # condition-inputs are unknown for these sequence-only predictors —
            # they carry no pH/temp/concentration, unlike a kinetic observable.
            "condition_inputs": "unknown",
            "available": value is not None}


def _waltz_sub_axis(n_waltz_amyloid: int) -> dict:
    """Waltz sub-axis as COVERAGE, not magnitude (review item C6)."""
    return {
        "value": None,                 # COVERAGE, not a propensity value
        "granularity": "hexapeptide",
        "endpoint": "amyloid (hexapeptide)",
        "reduction": "coverage: count of Waltz-validated amyloid-positive hexapeptides",
        "version": PREDICTOR_VERSION,
        "summary": None,
        "condition_inputs": "unknown",
        "note": ("Waltz here is COVERAGE (validated-hexapeptide count), NOT a "
                 "propensity magnitude; absence = UNTESTED, not non-amyloid"),
        "n_waltz_amyloid_peptides": int(n_waltz_amyloid),
        # 'available' reflects coverage: do we have a Waltz-tested amyloid hexapeptide?
        "available": n_waltz_amyloid > 0,
    }


def _no_wt_predictors() -> dict:
    """Predictor block for a protein with only engineered-mutant peptides."""
    reason = "only mutant peptides deposited"
    base = {"value": None, "version": PREDICTOR_VERSION, "summary": None,
            "condition_inputs": "unknown", "available": False, "reason": reason}
    return {
        "tango": {**base, "granularity": "residue/segment",
                  "endpoint": "beta-aggregation nucleation",
                  "reduction": "max segment TANGO score over WT peptides"},
        "aggrescan": {**base, "granularity": "residue",
                      "endpoint": "generic-aggregation",
                      "reduction": "max normalized aggregation propensity over WT peptides"},
        "pasta": {**base, "granularity": "residue-pair",
                  "endpoint": "amyloid (cross-beta)",
                  "reduction": "min (most stable) PASTA pairing energy over WT peptides"},
        "waltz": {**base, "granularity": "hexapeptide",
                  "endpoint": "amyloid (hexapeptide)",
                  "reduction": "coverage: count of Waltz-validated amyloid hexapeptides",
                  "note": "Waltz here is COVERAGE; absence = UNTESTED, not non-amyloid",
                  "n_waltz_amyloid_peptides": 0},
    }


# ------------------------------- APR sheet --------------------------------- #
def parse_aprs(path: Path) -> dict:
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb["final data"]
    h = _header_index(ws)
    iu = h.get("Uniprot ID")
    by_up: dict[str, list] = defaultdict(list)
    for r in ws.iter_rows(min_row=2, values_only=True):
        up = _clean(r[iu]) if iu is not None else None
        if not up:
            continue
        by_up[up].append({
            "position": _clean(r[h["Sequence Position"]]) if "Sequence Position" in h else None,
            "length": _num(r[h["APR Length"]]) if "APR Length" in h else None,
            "region": _clean(r[h["Experimental Aggregating Region"]]) if "Experimental Aggregating Region" in h else None,
            "category": _clean(r[h["Category"]]) if "Category" in h else None,
        })
    return dict(by_up)


# ---------------------------- structure sheet ------------------------------ #
def parse_structures(path: Path) -> dict:
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb["final data"]
    h = _header_index(ws)
    iu = h.get("UniProt AC")
    by_up: dict[str, list] = defaultdict(list)
    for r in ws.iter_rows(min_row=2, values_only=True):
        up = _clean(r[iu]) if iu is not None else None
        if not up:
            continue
        by_up[up].append({
            "pdb_id": _clean(r[h["PDB-ID"]]) if "PDB-ID" in h else None,
            "amyloid": _clean(r[h["amyloid/Non-amyloid"]]) if "amyloid/Non-amyloid" in h else None,
            "type": _clean(r[h["Type"]]) if "Type" in h else None,
            "method": _clean(r[h["Method"]]) if "Method" in h else None,
            "resolution": _num(r[h["Resolution"]]) if "Resolution" in h else None,
        })
    return dict(by_up)


# ------------------------------- assembly ---------------------------------- #
def run_etl(rawdir: Path, outdir: Path) -> dict:
    pep_path = rawdir / "aggregating peptides.xlsx"
    apr_path = rawdir / "APR information.xlsx"
    str_path = rawdir / "amyloid structure.xlsx"

    predictors = parse_peptides(pep_path)
    aprs = parse_aprs(apr_path)
    structures = parse_structures(str_path)

    all_ups = set(predictors) | set(aprs) | set(structures)
    by_uniprot = {}
    for up in sorted(all_ups):
        rec = predictors.get(up, {"protein_name": None, "n_peptides": 0,
                                   "n_amyloid_peptides": 0, "predictors": None})
        by_uniprot[up] = {
            "uniprot_id": up,
            "protein_name": rec.get("protein_name"),
            "n_peptides": rec.get("n_peptides", 0),
            "n_peptides_total": rec.get("n_peptides_total", rec.get("n_peptides", 0)),
            "n_peptides_wt": rec.get("n_peptides_wt", 0),
            "n_amyloid_peptides": rec.get("n_amyloid_peptides", 0),
            "peptide_length_range": rec.get("peptide_length_range"),
            "excluded_available_predictors": rec.get("excluded_available_predictors", []),
            "sequence_predictors": rec.get("predictors"),
            "aprs": aprs.get(up, []),
            "structures": structures.get(up, []),
        }

    outdir.mkdir(parents=True, exist_ok=True)
    payload = {"schema_version": SCHEMA_VERSION,
               "absent_predictors": {
                   "zyggregator": "not present in CPAD 2.0 bundle (declared absent)",
                   "camsol": "not present in CPAD 2.0 bundle (declared absent; "
                             "CamSol is a SOLUBILITY endpoint — distinct from the others)"},
               "by_uniprot": by_uniprot}
    with open(outdir / "sequence_structure.json", "w", encoding="utf-8") as fh:
        json.dump(payload, fh, indent=2)

    def _has_available_predictor(v):
        sp = v.get("sequence_predictors") or {}
        return any((sub or {}).get("available") for sub in sp.values())

    manifest = {
        "schema_version": SCHEMA_VERSION,
        "predictor_version_tag": PREDICTOR_VERSION,
        "n_uniprots": len(by_uniprot),
        # n_with_predictors now counts proteins that have >=1 AVAILABLE predictor
        # (zero-WT proteins emit a predictor block but with everything unavailable).
        "n_with_predictors": sum(1 for v in by_uniprot.values() if _has_available_predictor(v)),
        "n_with_wt_peptides": sum(1 for v in by_uniprot.values() if v["n_peptides_wt"]),
        "n_with_aprs": sum(1 for v in by_uniprot.values() if v["aprs"]),
        "n_with_structures": sum(1 for v in by_uniprot.values() if v["structures"]),
        "sources": {
            "peptides": {"file": pep_path.name, "sha256": sha256_of(pep_path)},
            "aprs": {"file": apr_path.name, "sha256": sha256_of(apr_path)},
            "structures": {"file": str_path.name, "sha256": sha256_of(str_path)},
        },
    }
    with open(outdir / "sequence_manifest.json", "w", encoding="utf-8") as fh:
        json.dump(manifest, fh, indent=2)
    return manifest


def main(argv=None):
    root = Path(__file__).resolve().parent.parent
    ap = argparse.ArgumentParser(description="PRISE CPAD sequence/structure ETL")
    ap.add_argument("--rawdir", type=Path, default=root / "data" / "cpad_raw")
    ap.add_argument("--outdir", type=Path, default=root / "data" / "processed")
    args = ap.parse_args(argv)
    for f in ("aggregating peptides.xlsx", "APR information.xlsx",
              "amyloid structure.xlsx"):
        if not (args.rawdir / f).exists():
            raise SystemExit(f"missing source: {args.rawdir / f}")
    manifest = run_etl(args.rawdir, args.outdir)
    print(f"[seq-etl] wrote sequence_structure.json + manifest to {args.outdir}")
    print(json.dumps(manifest, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
